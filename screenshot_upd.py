import sys
import os
import pyperclip
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image, ImageDraw, ImageFont
from openpyxl.styles import Font, Alignment


def text_to_image(text, img_path):
    """Convert text into an image with Consolas font, green text, and black background, then resize to 75%."""
    font_path = "consola.ttf"  # Ensure Consolas is available
    font = ImageFont.truetype(font_path, 24)
    padding = 20  # More padding for better spacing

    # Estimate text size dynamically
    img_temp = Image.new("RGB", (1000, 500), "black")  # Large temp image
    draw_temp = ImageDraw.Draw(img_temp)

    # Use textbbox to get precise text dimensions
    bbox = draw_temp.textbbox((0, 0), text, font=font)
    text_width = bbox[2] - bbox[0]  # Right - Left
    text_height = bbox[3] - bbox[1]  # Bottom - Top

    # Adjust final image size based on text dimensions
    width = text_width + 2 * padding
    height = text_height + 2 * padding

    # Create final image with exact size
    img = Image.new("RGB", (width, height), "black")
    draw = ImageDraw.Draw(img)

    # Draw text centered with padding
    draw.text((padding, padding), text, font=font, fill="green")

    # Resize to 75% while keeping proportions
    new_size = (int(width * 0.75), int(height * 0.75))
    img = img.resize(new_size, Image.Resampling.LANCZOS)

    # Save the final image
    img.save(img_path)


def write_text_to_excel(excel_file, sheet_name, row, col, text, is_clipboard):
    # Check if the Excel file exists
    if os.path.exists(excel_file):
        wb = load_workbook(excel_file)
        if sheet_name not in wb.sheetnames:
            wb.create_sheet(sheet_name)
        sheet = wb[sheet_name]
    else:
        wb = Workbook()
        sheet = wb.active
        sheet.title = sheet_name

    # Write the text into the specified cell
    cell = sheet[f"{col}{row}"]
    cell.value = text

    # If the text is from clipboard, adjust column width and font
    if is_clipboard:
        # Set the column width to 105
        sheet.column_dimensions[col].width = 105

        # Set the row height to 389
        sheet.row_dimensions[int(row)].height = 389

        # Set the font to Consolas
        cell.font = Font(name='Consolas', size=12)


    # Enable text wrapping
    cell.alignment = Alignment(wrap_text=True)


    # Save the Excel file
    try:
        wb.save(excel_file)
        print(f"Successfully saved the file: {excel_file}")
    except PermissionError:
        print(f"Error: Permission denied. Please ensure that '{excel_file}' is closed.")
        sys.exit(1)

def write_image_to_excel(excel_file, sheet_name, row, col, img_path):
    """Writes an image into an Excel file."""
    if os.path.exists(excel_file):
        wb = load_workbook(excel_file)
        if sheet_name not in wb.sheetnames:
            wb.create_sheet(sheet_name)
        sheet = wb[sheet_name]
    else:
        wb = Workbook()
        sheet = wb.active
        sheet.title = sheet_name

    # Insert image into Excel
    img_obj = ExcelImage(img_path)
    sheet.add_image(img_obj, f"{col}{row}")

    wb.save(excel_file)
    print(f"Successfully saved image to {excel_file}")


if __name__ == "__main__":
    if len(sys.argv) < 6:
        print("Usage: python script.py <excel_file> <sheet_name> <row> <col> <CLIPBOARD|IMAGE>")
        sys.exit(1)

    excel_file = sys.argv[1]
    sheet_name = sys.argv[2]
    row = sys.argv[3]
    col = sys.argv[4]
    mode = sys.argv[5].upper()

    img_path = "clipboard_output.png"  # Temporary file for storing image

    if mode == "CLIPBOARD":
        try:
            text = pyperclip.paste()
            is_clipboard = True
        except Exception as e:
            print(f"Error accessing clipboard: {e}")
            sys.exit(1)

        write_text_to_excel(excel_file, sheet_name, row, col, text, is_clipboard)

    elif mode == "IMAGE":
        text = pyperclip.paste()
        if not text.strip():
            print("Error: No text found in clipboard.")
            sys.exit(1)

        text_to_image(text, img_path)  # Convert text to image
        write_image_to_excel(excel_file, sheet_name, row, col, img_path)

    else:
        # Join all arguments from the 5th onwards to preserve spaces
        text = " ".join(sys.argv[5:])
        is_clipboard = False
        write_text_to_excel(excel_file, sheet_name, row, col, text, is_clipboard)